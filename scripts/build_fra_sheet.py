"""Build a Reality Fracture (FRA) Limited ratings sheet from J2SJosh's MTG Arena Zone set review.

Ratings and commentary come from the seven review articles. Card metadata is taken from
Scryfall and cross-checked against the printed mana cost, not from the article text.
"""

from __future__ import annotations

import csv
import html
import json
import re
import sys
import time
import urllib.request
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
CACHE = ROOT / ".cache"
OUT = ROOT / "output"
UA = "fra-limited-sheet/1.0 (personal research script)"

ARTICLES = {
    "White": "https://mtgazone.com/reality-fracture-fra-limited-set-review-white/",
    "Blue": "https://mtgazone.com/reality-fracture-fra-limited-set-review-blue/",
    "Black": "https://mtgazone.com/reality-fracture-fra-limited-set-review-black/",
    "Red": "https://mtgazone.com/reality-fracture-fra-limited-set-review-red/",
    "Green": "https://mtgazone.com/reality-fracture-fra-limited-set-review-green/",
    "Artifacts and Lands": "https://mtgazone.com/reality-fracture-fra-limited-set-review-artifacts-and-lands/",
    "Multicolor": "https://mtgazone.com/reality-fracture-fra-limited-set-review-multicolor/",
}

# Headings that end the card-by-card portion of an article.
STOP_HEADINGS = {"Wrap Up", "Related Guides", "Popular on MTG Arena Zone"}

COLOR_NAMES = {"W": "White", "U": "Blue", "B": "Black", "R": "Red", "G": "Green"}
WUBRG = "WUBRG"

TOKEN_RE = re.compile(
    r"<h([23])[^>]*>(?P<heading>.*?)</h\1>"
    r"|<a target=\"_blank\" href=\"https://mtgazone\.com/cards/(?P<slug>[^\"/]+)/\" class=\"cardblock\">"
    r"\s*<div class=\"front\">\s*<img[^>]*data-src=\"(?P<img>[^\"]+)\""
    r"|<p class=\"wp-block-paragraph\">(?P<para>.*?)</p>"
    r"|<li>(?P<li>.*?)</li>",
    re.S,
)
RATING_RE = re.compile(r"^Rating:\s*(\d(?:\.\d+)?)\s*/\s*5\s*$")
SCRYFALL_ID_RE = re.compile(r"/([0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12})\.jpg")


def fetch(url: str, dest: Path, refresh: bool = False) -> bytes:
    if dest.exists() and not refresh:
        return dest.read_bytes()
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "*/*"})
    for attempt in range(4):
        try:
            data = urllib.request.urlopen(req, timeout=60).read()
            break
        except Exception:
            if attempt == 3:
                raise
            time.sleep(2 ** (attempt + 2))
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(data)
    return data


def text(fragment: str) -> str:
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", "", fragment))).strip()


def load_scryfall(refresh: bool) -> list[dict]:
    dest = CACHE / "scryfall_fra.json"
    if dest.exists() and not refresh:
        return json.loads(dest.read_text())
    url = (
        "https://api.scryfall.com/cards/search?q=e%3Afra&unique=prints"
        "&include_extras=true&include_variations=true&order=set"
    )
    cards = []
    while url:
        req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "application/json"})
        page = json.load(urllib.request.urlopen(req, timeout=60))
        cards += page["data"]
        url = page.get("next_page")
        time.sleep(0.12)
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(json.dumps(cards))
    return cards


def parse_article(section: str, url: str, refresh: bool) -> tuple[list[dict], list[str]]:
    """Return rated entries (each may cover several cards) and parse warnings."""
    slug = url.rstrip("/").rsplit("/", 1)[1]
    page = fetch(url, CACHE / "articles" / f"{slug}.html", refresh).decode("utf-8")
    body_start = page.find("Here’s the usual grading scale")
    if body_start < 0:
        body_start = page.find("grading scale")
    page = page[body_start:]

    entries: list[dict] = []
    warnings: list[str] = []
    seg: dict | None = None

    def close(seg):
        if seg is None:
            return
        if seg["cards"] and seg["rating"] is not None:
            entries.append(seg)
        elif seg["cards"]:
            warnings.append(f"{section}: '{seg['heading']}' has card(s) but no rating")
        elif seg["rating"] is not None:
            warnings.append(f"{section}: '{seg['heading']}' has a rating but no card image")

    for m in TOKEN_RE.finditer(page):
        if m.group("heading") is not None:
            heading = text(m.group("heading"))
            rm = RATING_RE.match(heading)
            if rm:
                if seg is None:
                    warnings.append(f"{section}: rating '{heading}' before any heading")
                    continue
                if seg["rating"] is not None:
                    warnings.append(f"{section}: second rating under '{seg['heading']}'")
                seg["rating"] = rm.group(1)
                seg["rating_raw"] = heading
                continue
            if heading.lower().startswith("rating"):
                warnings.append(f"{section}: unparseable rating heading '{heading}'")
            close(seg)
            if heading in STOP_HEADINGS:
                seg = None
                break
            seg = {"heading": heading, "cards": [], "rating": None, "rating_raw": None,
                   "pre": [], "post": [], "section": section, "url": url}
        elif seg is None:
            continue
        elif m.group("slug"):
            seg["cards"].append({"slug": m.group("slug"), "img": html.unescape(m.group("img"))})
        else:
            frag = m.group("para") if m.group("para") is not None else m.group("li")
            t = text(frag)
            if t:
                (seg["post"] if seg["rating"] is not None else seg["pre"]).append(t)
    close(seg)
    return entries, warnings


def mana_symbols(cost: str) -> list[str]:
    return re.findall(r"\{([^}]+)\}", cost or "")


def mana_value(cost: str) -> float:
    """Mana value per CR 202.3: X is 0, hybrid counts its largest component, Phyrexian counts 1."""
    total = 0.0
    for sym in mana_symbols(cost):
        parts = sym.split("/")
        if sym in ("X", "Y", "Z"):
            continue
        if sym.isdigit():
            total += int(sym)
        elif sym == "½" or sym == "H":
            total += 0.5
        elif len(parts) == 2 and parts[0].isdigit():
            total += max(int(parts[0]), 1)
        else:
            total += 1
    return total


def colors_from_cost(cost: str) -> list[str]:
    found = {c for sym in mana_symbols(cost) for c in sym.split("/") if c in WUBRG}
    return [c for c in WUBRG if c in found]


def color_label(colors: list[str]) -> str:
    colors = [c for c in WUBRG if c in colors]
    if not colors:
        return "Colorless"
    if len(colors) == 1:
        return COLOR_NAMES[colors[0]]
    return "Multicolor (" + "/".join(COLOR_NAMES[c] for c in colors) + ")"


def fmt_num(v) -> str:
    if v is None or v == "":
        return ""
    f = float(v)
    return str(int(f)) if f.is_integer() else str(f)


def main() -> int:
    refresh = "--refresh" in sys.argv
    embed = "--no-images" not in sys.argv

    prints = load_scryfall(refresh)
    by_id = {c["id"]: c for c in prints}
    by_name: dict[str, list[dict]] = {}
    for c in prints:
        by_name.setdefault(c["name"].lower(), []).append(c)
        by_name.setdefault(c["name"].split(" // ")[0].lower(), []).append(c)

    def canonical(card: dict) -> dict:
        """Main-set printing of the same card (lowest collector number with the same oracle id)."""
        same = [c for c in prints if c["oracle_id"] == card["oracle_id"]]
        booster = [c for c in same if c.get("booster")] or same

        def key(c):
            n = re.match(r"\d+", c["collector_number"])
            return (int(n.group()) if n else 10**6, c["collector_number"])

        return min(booster, key=key)

    rows: list[dict] = []
    warnings: list[str] = []
    seen: dict[str, dict] = {}
    duplicates: list[str] = []
    notes: list[str] = []

    for section, url in ARTICLES.items():
        entries, warn = parse_article(section, url, refresh)
        warnings += warn
        for e in entries:
            group = len(e["cards"]) > 1
            for card_ref in e["cards"]:
                idm = SCRYFALL_ID_RE.search(card_ref["img"])
                sc = by_id.get(idm.group(1)) if idm else None
                if sc is None:
                    guess = card_ref["slug"].replace("-", " ")
                    cands = [c for k, v in by_name.items() if k.replace("'", "").replace(",", "").replace("//", "").replace("  ", " ") == guess for c in v]
                    sc = cands[0] if cands else None
                    if sc is None:
                        warnings.append(f"{section}: could not match '{card_ref['slug']}' to Scryfall")
                        continue
                    notes.append(f"{sc['name']}: article image is a non-FRA printing; matched to the FRA printing by name")
                sc = canonical(sc)
                front_name = sc["name"].split(" // ")[0]
                if not group and text(e["heading"]).replace("’", "'") != front_name.replace("’", "'"):
                    warnings.append(f"{section}: heading '{e['heading']}' vs Scryfall name '{sc['name']}'")

                if sc["oracle_id"] in seen:
                    prev = seen[sc["oracle_id"]]
                    duplicates.append(
                        f"{sc['name']}: reviewed in {prev['Source Section']} ({prev['J2SJosh Rating']}) "
                        f"and {section} ({e['rating']}); kept the first"
                    )
                    continue

                faces = sc.get("card_faces") or []
                main_face = faces[0] if faces and sc["layout"] != "normal" else sc
                cost = main_face.get("mana_cost", sc.get("mana_cost", ""))
                second = faces[1] if len(faces) > 1 else None

                mv_calc = mana_value(cost)
                if abs(mv_calc - float(sc["cmc"])) > 1e-9:
                    warnings.append(f"{sc['name']}: computed MV {mv_calc} != Scryfall cmc {sc['cmc']}")
                colors = sc.get("colors") if sc.get("colors") is not None else main_face.get("colors", [])
                colors = [c for c in WUBRG if c in colors]
                cost_colors = colors_from_cost(cost)
                if colors != cost_colors and not sc.get("color_indicator") and "Land" not in sc["type_line"]:
                    notes.append(f"{sc['name']}: color {colors} differs from mana cost colors {cost_colors}")

                img = sc.get("image_uris") or (faces[0].get("image_uris") if faces else {}) or {}
                commentary = e["post"] or e["pre"]
                if group:
                    commentary = e["pre"] + e["post"]
                row = {
                    "Card Name": sc["name"],
                    "Front Face Name": front_name,
                    "Color": color_label(colors),
                    "Color Code": "".join(colors) or "C",
                    "Color Identity": "".join(c for c in WUBRG if c in sc.get("color_identity", [])) or "C",
                    "Mana Cost": cost,
                    "Mana Value": fmt_num(sc["cmc"]),
                    "Rarity": sc["rarity"].capitalize(),
                    "Card Type": main_face.get("type_line", sc["type_line"]),
                    "Power": main_face.get("power", sc.get("power", "")),
                    "Toughness": main_face.get("toughness", sc.get("toughness", "")),
                    "Loyalty": main_face.get("loyalty", sc.get("loyalty", "")),
                    "Layout": sc["layout"],
                    "Second Half": second["name"] if second else "",
                    "Second Half Cost": second.get("mana_cost", "") if second else "",
                    "Second Half Type": second.get("type_line", "") if second else "",
                    "J2SJosh Rating": e["rating"],
                    "Rating Note": f"Group rating for '{e['heading']}'" if group else "",
                    "J2SJosh Commentary": "\n\n".join(commentary),
                    "Source Section": section,
                    "Source Article": url,
                    "Collector Number": sc["collector_number"],
                    "Scryfall URL": sc["scryfall_uri"].split("?")[0],
                    "Image URL": img.get("normal") or img.get("large") or "",
                    "Scryfall ID": sc["id"],
                    "_small": img.get("small", ""),
                }
                seen[sc["oracle_id"]] = row
                rows.append(row)

    rows.sort(key=lambda r: (int(re.match(r"\d+", r["Collector Number"]).group()), r["Collector Number"]))

    OUT.mkdir(exist_ok=True)
    cols = [k for k in rows[0] if not k.startswith("_")]
    with open(OUT / "fra_j2sjosh_limited_ratings.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)

    simple_cols = ["Card Name", "Color", "Mana Cost", "Mana Value", "Rarity", "Card Type",
                   "Power", "Toughness", "J2SJosh Rating", "Scryfall URL", "Image URL"]
    with open(OUT / "fra_j2sjosh_limited_ratings_simple.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=simple_cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(sorted(rows, key=lambda r: (-float(r["J2SJosh Rating"]), r["Card Name"])))

    write_xlsx(rows, cols, embed)

    all_cards = {c["oracle_id"] for c in prints}
    unreviewed = sorted({c["name"] for c in prints if c["oracle_id"] not in seen})
    report = {
        "unique_cards": len(rows),
        "cards_with_rating": sum(1 for r in rows if r["J2SJosh Rating"] != ""),
        "by_section": {s: sum(1 for r in rows if r["Source Section"] == s) for s in ARTICLES},
        "by_rarity": {k: sum(1 for r in rows if r["Rarity"] == k) for k in ["Common", "Uncommon", "Rare", "Mythic"]},
        "group_rated_cards": [r["Card Name"] + " (" + r["Rating Note"] + ")" for r in rows if r["Rating Note"]],
        "duplicates_skipped": duplicates,
        "parse_or_verification_warnings": warnings,
        "notes": notes,
        "unique_cards_in_set_on_scryfall": len(all_cards),
        "cards_in_set_not_reviewed": unreviewed,
    }
    (OUT / "verification_report.json").write_text(json.dumps(report, indent=2, ensure_ascii=False))
    print(json.dumps({k: v for k, v in report.items() if k != "group_rated_cards"}, indent=2, ensure_ascii=False))
    return 0


def write_xlsx(rows: list[dict], cols: list[str], embed: bool) -> None:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "FRA Ratings"
    headers = ["Card Image"] + cols
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    widths = {"Card Image": 18, "Card Name": 30, "J2SJosh Commentary": 90, "Card Type": 28,
              "Source Article": 30, "Scryfall URL": 30, "Image URL": 30, "Color": 22}
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 12)

    for r_i, row in enumerate(rows, 2):
        for c_i, h in enumerate(cols, 2):
            v = row[h]
            if h in ("J2SJosh Rating", "Mana Value") and v != "":
                v = float(v)
            cell = ws.cell(row=r_i, column=c_i, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=h == "J2SJosh Commentary")
            if h in ("Scryfall URL", "Image URL", "Source Article") and v:
                cell.hyperlink = v
                cell.font = Font(color="2563EB", underline="single")
        if embed and row["_small"]:
            data = fetch(row["_small"], CACHE / "images" / f"{row['Scryfall ID']}.jpg")
            img = XLImage(BytesIO(data))
            img.width, img.height = 122, 170
            ws.add_image(img, f"A{r_i}")
            ws.row_dimensions[r_i].height = 132
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    wb.save(OUT / "fra_j2sjosh_limited_ratings.xlsx")


if __name__ == "__main__":
    raise SystemExit(main())
